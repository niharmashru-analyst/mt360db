"""Builds sample/review_data.xlsx (same 4 sheets + headers as the real extract) and sample/masters.xlsx."""
import random, math, datetime as dt, openpyxl
from openpyxl import Workbook
random.seed(11)
OUT='sample/'
months=[(2025+(3+i)//12,(3+i)%12+1) for i in range(18)]   # Apr-25 .. Sep-26
fy=lambda y,m: f'{y}-{str(y+1)[2:]}' if m>=4 else f'{y-1}-{str(y)[2:]}'

# ---- SKU master (synthetic) -----------------------------------------------------------------
CATS={('Makeup','Lipstick'):['MATTE LIP COLOR {c} 4.2ML','LIQUID LIPSTICK {c} 5ML'],('Makeup','Eyeliner'):['MIDNIGHT EYELINER {c} 2.5ML','KAJAL {c} 0.35GM'],
('Makeup','Foundation'):['BB CREAM {c} 30ML','COMPACT POWDER {c} 9GM'],('Makeup','Brushes'):['BLENDING BRUSH R{n}','POWDER BRUSH R{n}'],
('Skincare','Sunscreen'):['SUNSCREEN GEL SPF 50 {n}0ML'],('Skincare','Serum'):['VITAMIN C SERUM {n}0ML'],('Skincare','Cleanser'):['FACE WASH {c} {n}0ML'],
('Fragrance','EDP'):['PERFUME {c} 50ML'],('Fragrance','Body Mist'):['BODY MIST {c} 100ML'],('Haircare','Hair Oil'):['HAIR OIL {c} 100ML']}
SH=['ROSE','PEACH','CORAL','MOCHA','NUDE','BERRY','HAZELNUT','OCEAN','VANILLA','MIDNIGHT']
skus=[];ean=8906121640000
for (cat,sub),tpls in CATS.items():
    for t in tpls:
        for i in range(random.choice([2,3,3])):
            ean+=random.randint(7,90)
            name=t.format(c=SH[len(skus)%len(SH)],n=random.randint(1,9))
            mrp=random.choice([299,349,399,449,499,550,599,650,699,799,999,1199])
            skus.append(dict(ean=ean,name=name,cat=cat,sub=sub,mrp=mrp,vel=random.choice([8,6,4,3,2,1.5,1,.6]),launch=''))
for i,s in enumerate(skus):
    if i%7==3: y,m=random.choice(months[3:-2]); s['launch']=f'{y}-{m:02d}'
print(len(skus),'skus')

# ---- Outlets ---------------------------------------------------------------------------------
CITY={'Chennai':('Tamil Nadu','South'),'Hyderabad':('Telangana','South'),'Bangalore':('Karnataka','South'),'Mumbai':('Maharashtra','West'),'Pune':('Maharashtra','West'),
'Ahmedabad':('Gujarat','West'),'Surat':('Gujarat','West'),'New Delhi':('Delhi','North'),'Noida':('Uttar Pradesh','North'),'Lucknow':('Uttar Pradesh','North'),
'Chandigarh':('Punjab','North'),'Jaipur':('Rajasthan','North'),'Kolkata':('West Bengal','East'),'Bhubaneswar':('Odisha','East'),'Guwahati':('Assam','East'),'Kochi':('Kerala','South')}
BGR={'North':'BGR North','West':'BGR West','South':'BGR South','East':'BGR East'}
SHEETS={'Health & Glow':('MBO',26,'HG'),'Shoppers Stop':('Department Store',18,'SSL'),'Dabur':('MBO',22,'DAB'),'Lifestyle':('Department Store',20,'LS')}
outlets=[];n=1000
for sheet,(typ,cnt,pfx) in SHEETS.items():
    for i in range(cnt):
        city=random.choice(list(CITY));st,rg=CITY[city];n+=1
        outlets.append(dict(sheet=sheet,code=f'OL{n:04d}',site=random.randint(100,3999),name=f'{pfx}-{random.choice(["PHOENIX","CITY CENTRE","FORUM","ORION","EXPRESS","INORBIT","SELECT","NEXUS","GALLERIA"])} {city.upper()} {i+1}',
             city=city,state=st,region=rg,typ=typ,w=random.uniform(.4,1.8),trend=random.uniform(-.3,.45),mp=random.randint(1,4),vis=random.uniform(45,95),null_code=(random.random()<.12)))

# ---- transactions ----------------------------------------------------------------------------
def lines(o,y,m,mi):
    out=[]
    seas=1+.3*math.cos((m-11)/12*2*math.pi)
    for s in skus:
        if s['launch'] and f'{y}-{m:02d}'<s['launch']: continue
        if random.random()>0.55: continue
        mu=s['vel']*o['w']*seas*(1+o['trend']*mi/17)*random.lognormvariate(0,.3)
        q=int(round(mu))
        if q<=0 and random.random()<.7: continue
        k=1 if o['sheet'] in('Shoppers Stop',) else random.randint(1,4)   # transaction lines
        for j in range(k):
            ql=max(1,round(q/k)) if q>0 else 1
            if random.random()<.03: ql=-1                                  # returns
            out.append((s,ql,random.randint(1,28)))
    return out

wb=Workbook(write_only=True)
HEAD={'Health & Glow':['Year','Month','Location Code','Location Name','Outlet Code','Short Name','Sku Name','Ean Code','Sales Qty','MRP','Sales on MRP'],
'Shoppers Stop':['FY','Mnth & Year','Region','Outlet Code','Site Name','City','EAN','Description','MRP','Site','Qty in unit of entry','Sales Value inc. VAT'],
'Dabur':['FY','Mnth & Year','Posting Date','Site','OL Code','Site Name','Description','EAN','Qty in unit of entry','Sales Value inc. VAT','MRP'],
'Lifestyle':['FY Year','SALEDATE','Month','LOC_NUM','Ol Code','LOCATION_NAME','City','EAN_CODE','ITEM_DESC_SECONDARY','SHORT_DESC','SUP_NAME','QTY','SALRRP','SALNOT','SALMRP','Offer']}
sales_val={}   # (month, outlet, ean) -> qty for stock/target generation
for sheet in SHEETS:
    ws=wb.create_sheet(sheet);ws.append(HEAD[sheet])
    for o in [x for x in outlets if x['sheet']==sheet]:
        for mi,(y,m) in enumerate(months):
            d0=dt.datetime(y,m,1)
            for s,q,day in lines(o,y,m,mi):
                mrp=s['mrp']; code=None if (o['null_code'] and random.random()<.7) else o['code']
                sales_val[(f'{y}-{m:02d}',o['code'],s['ean'])]=sales_val.get((f'{y}-{m:02d}',o['code'],s['ean']),0)+q
                dd=dt.datetime(y,m,day)
                if sheet=='Health & Glow':
                    mr=0 if random.random()<.01 else mrp
                    ws.append([fy(y,m),d0,o['site'],o['name'],code,'RENEE COS',s['name'],s['ean'],0 if random.random()<.05 else q,mr,q*mr])
                elif sheet=='Shoppers Stop':
                    ws.append([fy(y,m),d0,o['region'],code,o['name'],o['city'],s['ean'],s['name'],mrp,o['site'],q,q*mrp])
                elif sheet=='Dabur':
                    ws.append([fy(y,m),d0,dd,o['site'],code,o['name'],'Renee '+s['name'].title(),s['ean'],q,q*mrp,mrp])
                else:
                    net=round(mrp*random.uniform(.78,.9),2)
                    ws.append([fy(y,m),dd,f"{dt.date(y,m,1).strftime('%B')}'{str(y)[2:]}",o['site'],code,o['name'],o['city'],s['ean'],s['name'],s['ean'],'RENEE COSMETICS PRIVATE LIMITED',q,round(net*1.18,2),net*q,mrp,None])
wb.save(OUT+'review_data.xlsx'); print('review_data.xlsx saved')

# ---- masters.xlsx ----------------------------------------------------------------------------
mw=Workbook(write_only=True)
w=mw.create_sheet('Outlet Master');w.append(['Outlet Code','Outlet Name','Chain Name','Chain Type','City','State','Region','BGR','Manpower','Visibility %'])
for o in outlets: w.append([o['code'],o['name'],o['sheet'],o['typ'],o['city'],o['state'],o['region'],BGR[o['region']],o['mp'],round(o['vis'],1)])
w=mw.create_sheet('SKU Master');w.append(['EAN','SKU Code','SKU','Brand','Category','Sub Category','Status','MRP','Launch Date','Margin %'])
MG={'Makeup':50,'Skincare':45,'Fragrance':55,'Haircare':35}
for s in skus: w.append([s['ean'],f"SKU{s['ean']%100000:05d}",s['name'].title(),'Renee',s['cat'],s['sub'],'Active',s['mrp'],s['launch'],MG[s['cat']]])
w=mw.create_sheet('Targets');w.append(['Month','Outlet Code','Target Value'])
tot={}
for (mo,oc,e),q in sales_val.items(): tot[(mo,oc)]=tot.get((mo,oc),0)+max(q,0)*next(s['mrp'] for s in skus if s['ean']==e)
for (mo,oc),v in tot.items(): w.append([mo,oc,round(v*random.uniform(.9,1.15)*1.05)])
w=mw.create_sheet('Stock & Orders');w.append(['Month','Outlet Code','EAN','OP Stock','CL Stock','Order Qty','Filled Qty'])
prev={}
for (mo,oc,e),q in sorted(sales_val.items()):
    q=max(q,0);op=prev.get((oc,e),q*random.randint(1,3));od=int(q*random.uniform(.8,2))+random.randint(0,3);fl=int(od*random.choice([1,1,.95,.9,.8,.7]))
    r=random.random();cl=0 if r<.07 else (int(q*random.uniform(5,9)) if r<.12 else max(0,op+fl-q));prev[(oc,e)]=cl
    w.append([mo,oc,e,op,cl,od,fl])
mw.save(OUT+'masters.xlsx');print('masters.xlsx saved',len(outlets),'outlets')
